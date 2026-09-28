# =============================================================================
# report.py  — Word (.docx) report generator
# =============================================================================
import io
import datetime
import textwrap
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from .constants import TARGET, CHART_DIR, REPORT_PATH
from .charts import cm_chart


# ── helpers ──────────────────────────────────────────────────────────────────
def _h(doc, text, level=1):
    return doc.add_heading(text, level=level)

def _p(doc, text, bold=False, italic=False, indent=False, size=11):
    p   = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold; run.italic = italic
    run.font.size = Pt(size)
    if indent:
        p.paragraph_format.left_indent = Pt(18)
    return p

def _tbl(doc, df_in: pd.DataFrame, title=""):
    if title:
        _p(doc, title, bold=True)
    tbl = doc.add_table(rows=1, cols=len(df_in.columns))
    tbl.style = "Light Grid Accent 1"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0].cells
    for i, col in enumerate(df_in.columns):
        hdr[i].text = str(col)
        for r in hdr[i].paragraphs[0].runs:
            r.bold = True
    for _, row_data in df_in.iterrows():
        row = tbl.add_row().cells
        for i, val in enumerate(row_data):
            row[i].text = str(round(val, 4) if isinstance(val, float) else val)

def _code(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(12)
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(7.5)

def _pic(doc, path, width=5.5):
    if path and Path(path).exists():
        doc.add_picture(str(path), width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


# ── main generator ────────────────────────────────────────────────────────────
def generate_word_report(df: pd.DataFrame, artefacts: dict,
                          eda_chart_paths: dict) -> bytes:
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Inches(1)
        sec.left_margin = sec.right_margin = Inches(1.1)

    # Cover ───────────────────────────────────────────────────────────────────
    doc.add_paragraph(); doc.add_paragraph()
    tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r  = tp.add_run("Social Media Impact Prediction")
    r.bold = True; r.font.size = Pt(24); r.font.color.rgb = RGBColor(0x1A, 0x73, 0xE8)
    sp = doc.add_paragraph(); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sp.add_run("Using Machine Learning"); r2.bold = True; r2.font.size = Pt(18)
    doc.add_paragraph()
    for label, val in [
        ("Project Type",    "Machine Learning Classification"),
        ("Dataset",         "Social_media_impact_on_life.csv"),
        ("Target Variable", "Overall_Impact"),
        ("Technologies",    "Python · Streamlit · scikit-learn · pandas · python-docx"),
    ]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"{label}: ").bold = True
        p.add_run(val)
    doc.add_page_break()

    # 1. Abstract ─────────────────────────────────────────────────────────────
    _h(doc, "1. Abstract")
    _p(doc, ("This project investigates the relationship between social media usage patterns "
             "and their overall impact on students' well-being, academic performance, and mental "
             "health. Three supervised classifiers (Logistic Regression, Decision Tree, Random Forest) "
             "predict the Overall_Impact category (Beneficial / Neutral / Negative). The project "
             "delivers an interactive Streamlit multi-page app with EDA, model evaluation, real-time "
             "prediction, an AI wellbeing assistant (English & Hindi), and this auto-generated report."))

    # 2. Problem Statement ────────────────────────────────────────────────────
    _h(doc, "2. Problem Statement")
    _p(doc, ("Given measurable student indicators (usage hours, sleep, stress scores, GPA, etc.), "
             "can a ML model predict whether social media has a Beneficial, Neutral, or Negative "
             "overall impact? The answer has practical value for educators, counsellors, and students — "
             "though it must be understood as a statistical pattern, not a medical diagnosis."))

    # 3. Dataset ──────────────────────────────────────────────────────────────
    _h(doc, "3. Dataset Description")
    info = [("Dataset", "Social_media_impact_on_life.csv"),
            ("Records", str(len(df))), ("Columns", str(len(df.columns))),
            ("Target", TARGET), ("Classes", "Beneficial, Neutral, Negative"),
            ("Missing Values", "Perceived_Stress_Score: 46 | Academic_Performance_GPA: 85"),
            ("Duplicates", "0")]
    tbl = doc.add_table(rows=len(info)+1, cols=2)
    tbl.style = "Light Grid Accent 1"
    tbl.rows[0].cells[0].text = "Property"; tbl.rows[0].cells[1].text = "Value"
    for c in tbl.rows[0].cells:
        for run in c.paragraphs[0].runs: run.bold = True
    for i, (k, v) in enumerate(info, 1):
        tbl.rows[i].cells[0].text = k; tbl.rows[i].cells[1].text = v
    doc.add_paragraph()
    tc = df[TARGET].value_counts().reset_index()
    tc.columns = ["Category", "Count"]
    tc["Percentage"] = (tc["Count"]/len(df)*100).round(2).astype(str)+"%"
    _tbl(doc, tc, "Target Distribution")

    # 4. EDA ──────────────────────────────────────────────────────────────────
    _h(doc, "4. Exploratory Data Analysis")
    eda_items = [
        ("target_dist",     "Overall Impact Distribution",
         "Dataset is highly imbalanced: 81.8% Beneficial, 14.5% Neutral, 3.7% Negative."),
        ("daily_usage",     "Daily Usage Hours Distribution",
         "Most students use social media 3–9 hours daily."),
        ("sleep_dur",       "Sleep Duration Distribution",
         "Majority sleep 6–8 hours; less sleep trends toward higher stress."),
        ("stress",          "Perceived Stress Score Distribution",
         "Stress clusters at 15–22; higher scores correlate with Negative impact."),
        ("gpa",             "Academic Performance GPA Distribution",
         "GPA concentrated 3.0–4.0; slight left tail for lower performers."),
        ("usage_vs_impact", "Daily Usage Hours vs Overall Impact",
         "Negative-impact students show slightly higher median usage."),
        ("sleep_vs_impact", "Sleep Duration vs Overall Impact",
         "Better-rested students trend toward Beneficial impact."),
        ("stress_vs_impact","Perceived Stress vs Overall Impact",
         "Negative-impact students report notably higher stress."),
        ("mhi_vs_impact",   "Mental Health Index vs Overall Impact",
         "Higher MHI is clearly associated with Beneficial impact."),
        ("gpa_vs_impact",   "GPA vs Overall Impact",
         "Modest but consistent GPA differences across impact categories."),
        ("cat_features",    "Categorical Feature Distributions",
         "No single platform dominates the Negative category."),
        ("corr_heatmap",    "Correlation Heatmap",
         "MHI and Sleep_Quality positively correlate with GPA; Stress inversely."),
        ("lnu_vs_impact",   "Late Night Usage vs Overall Impact",
         "Late-night users appear more in Neutral and Negative categories."),
    ]
    for key, title, insight in eda_items:
        _h(doc, title, level=2)
        _pic(doc, eda_chart_paths.get(key))
        _p(doc, f"Insight: {insight}", italic=True)

    # 5. Preprocessing ────────────────────────────────────────────────────────
    _h(doc, "5. Data Preprocessing")
    for step, desc in [
        ("Missing Values", "Median imputation via SimpleImputer inside pipeline (train-only fit prevents leakage)."),
        ("Identifier Exclusion", "Student_ID dropped."),
        ("Boolean Encoding", "Late_Night_Usage → int."),
        ("Numerical Features", "9 columns — StandardScaler."),
        ("Categorical Features", "5 columns — OneHotEncoder(handle_unknown='ignore')."),
        ("Target Encoding", "Beneficial=0, Negative=1, Neutral=2."),
        ("Train/Test Split", "80/20 stratified."),
        ("Class Imbalance", "Balanced class weights on all classifiers."),
    ]:
        _p(doc, f"• {step}: ", bold=True)
        _p(doc, desc, indent=True)

    # 6. Models ───────────────────────────────────────────────────────────────
    _h(doc, "6. Machine Learning Models")
    for name, desc in [
        ("Logistic Regression", "Linear baseline. C=1.0, max_iter=1000, class_weight='balanced'."),
        ("Decision Tree",       "max_depth=8, min_samples_leaf=10, class_weight='balanced'."),
        ("Random Forest",       "n_estimators=200, max_depth=10, min_samples_leaf=5, class_weight='balanced'."),
    ]:
        _h(doc, name, level=2); _p(doc, desc)

    # 7. Evaluation ───────────────────────────────────────────────────────────
    _h(doc, "7. Model Evaluation")
    res_df = artefacts["results_df"].copy()
    for col in ["Accuracy", "Precision", "Recall", "F1-Score"]:
        res_df[col] = res_df[col].apply(lambda x: f"{float(x)*100:.2f}%")
    _tbl(doc, res_df, "Model Comparison (Test Set)")
    le_classes = list(artefacts["le"].classes_)
    for mname, mres in artefacts["results"].items():
        _h(doc, mname, level=2)
        _p(doc, (f"Accuracy: {mres['Accuracy']*100:.2f}%  |  "
                 f"Precision: {mres['Precision']*100:.2f}%  |  "
                 f"Recall: {mres['Recall']*100:.2f}%  |  "
                 f"F1: {mres['F1-Score']*100:.2f}%"))
        _p(doc, "Classification Report:", bold=True)
        _code(doc, mres["cr"])
        cm_buf = cm_chart(mres["cm"], le_classes, mname)
        cm_path = CHART_DIR / f"cm_{mname.replace(' ','_').lower()}.png"
        if cm_path.exists():
            _pic(doc, cm_path, width=4)
    _p(doc, f"\nBest Model: {artefacts['best_name']} (highest weighted F1).", bold=True)
    fi_path = CHART_DIR / "14_feature_importance.png"
    if fi_path.exists():
        _h(doc, "Feature Importance (Random Forest)", level=2)
        _pic(doc, fi_path)
        _p(doc, "Mental_Health_Index, Perceived_Stress_Score, Sleep_Duration_Hours, and GPA are most predictive.", italic=True)

    # 8. Prediction System ────────────────────────────────────────────────────
    _h(doc, "8. Prediction System")
    _p(doc, "Student inputs are assembled into a single-row DataFrame, passed through the trained pipeline, and classified by the best model.")
    _code(doc, textwrap.dedent("""
        def predict_impact(artefacts, input_dict):
            row = pd.DataFrame([input_dict])
            pred_enc = artefacts["best_model"].predict(row)[0]
            return artefacts["le"].inverse_transform([pred_enc])[0]
    """))

    # 9. AI Assistant ─────────────────────────────────────────────────────────
    _h(doc, "9. AI Assistant")
    _p(doc, ("The AI Assistant page reads the last prediction from session state and responds "
             "in English and Hindi/Hinglish. It explains prediction factors, gives personalised "
             "wellbeing suggestions, handles what-if scenarios, and strictly avoids medical diagnoses. "
             "All responses are grounded in the user's actual input data."))

    # 10. Source Code ─────────────────────────────────────────────────────────
    _h(doc, "10. Complete Source Code")
    try:
        base = Path(__file__).resolve().parent.parent
        files = [
            base / "SakshiSingh_SocialMediaImpactPredictionModel.py",
        ] + sorted((base / "pages").glob("*.py"))
        for fpath in files:
            if fpath.exists():
                _h(doc, fpath.name, level=2)
                src = fpath.read_text(encoding="utf-8").splitlines()
                for start in range(0, len(src), 80):
                    _code(doc, "\n".join(src[start:start+80]))
    except Exception as e:
        _code(doc, f"# Source code could not be embedded: {e}")

    # 11. Results ─────────────────────────────────────────────────────────────
    _h(doc, "11. Results and Insights")
    _p(doc, "Key Findings:", bold=True)
    for insight in [
        "Dataset heavily skewed toward 'Beneficial' (82%) — class-weight balancing was essential.",
        "MHI, stress, and sleep quality are most strongly differentiated across impact categories.",
        "Late-night usage is more prevalent among Neutral and Negative impact students.",
        "Platform choice alone is a weak differentiator — behavioural patterns matter more.",
    ]:
        _p(doc, f"• {insight}", indent=True)
    best_res = artefacts["results"][artefacts["best_name"]]
    _p(doc, (f"Best model ({artefacts['best_name']}): "
             f"{best_res['Accuracy']*100:.2f}% accuracy, "
             f"{best_res['F1-Score']*100:.2f}% weighted F1."), indent=True)

    # 12. Conclusion ──────────────────────────────────────────────────────────
    _h(doc, "12. Conclusion")
    _p(doc, (f"This project demonstrates a complete end-to-end ML pipeline. "
             f"{artefacts['best_name']} performed best. Key predictive features: "
             "Mental_Health_Index, Perceived_Stress_Score, Sleep_Duration_Hours, GPA. "
             "The multi-page Streamlit app includes an AI assistant supporting English and Hindi. "
             "All results are statistical correlations — not causal claims, not medical diagnoses."))

    # 13. Requirements ────────────────────────────────────────────────────────
    _h(doc, "13. Requirements and Execution")
    _code(doc, "pip install pandas numpy matplotlib seaborn scikit-learn streamlit python-docx")
    _code(doc, "streamlit run SakshiSingh_SocialMediaImpactPredictionModel.py")

    # Footer ──────────────────────────────────────────────────────────────────
    doc.add_page_break()
    fp = doc.add_paragraph(); fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r  = fp.add_run("Social Media Impact Prediction — Auto-generated Report")
    r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x88, 0x88, 0x88); r.italic = True

    buf = io.BytesIO(); doc.save(buf); buf.seek(0)
    return buf.read()
